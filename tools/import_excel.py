#!/usr/bin/env python3
"""Convertit le classeur Excel de suivi des congés en data/seed.js.

Usage : python3 tools/import_excel.py AF_Activites.xlsx [data/seed.js]

Un onglet par année (nom = l'année). Dans chaque onglet :
  - grille jours : 12 mois x 2 colonnes (date, code), lignes 3 à 33 ;
  - colonne AB : codes (AB3..) avec droits en AC, jours fériés (dates),
    lieux (cellules colorées) et vacances scolaires (dates en AD/AE).
"""
import calendar
import colorsys
import datetime as dt
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

SRC = sys.argv[1] if len(sys.argv) > 1 else "AF_Activites.xlsx"
DST = sys.argv[2] if len(sys.argv) > 2 else "data/seed.js"

CODE_ALIASES = {"CP": "CA"}  # congés payés (2018-2021) = congés annuels CA01
QUOTA_CODES = {"CA", "CJT", "CCA", "RA", "CH"}
KNOWN_CODES = {"A", "TT", "CA", "CP", "CJT", "CCA", "RA", "CH", "Dmgt", "ML", "F", "SP", "OFF"}
# Légende colorée devenue un type de jour : nom dans la légende -> code
PLACE_AS_TYPE = {"PI Event": "PI"}
NOT_A_PLACE = {"Type", "Max", "Droits", "Jours ouvrés", "Jours trav.", "Jours fériés",
               "Vacances scolaires", "Hiver", "Printemps", "Eté", "Été", "Toussaint", "Noël"} | KNOWN_CODES


def fill_sig(cell):
    if cell.fill is None or cell.fill.fill_type != "solid":
        return None
    c = cell.fill.fgColor
    sig = c.rgb if c.type == "rgb" else ("t", c.theme, round(c.tint or 0, 2))
    return None if sig in (("t", 0, 0.0), ("t", 0, -0.25)) else sig


def theme_colors(wb):
    """Couleurs du thème Office : index 0..9 = lt1, dk1, lt2, dk2, accent1..6."""
    xml = wb.loaded_theme.decode("utf-8") if isinstance(wb.loaded_theme, bytes) else wb.loaded_theme
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S).group(0)
    vals = {}
    for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"):
        m = re.search(r"<a:%s>.*?(?:val|lastClr)=\"([0-9A-Fa-f]{6})\"" % tag, scheme, re.S)
        vals[tag] = m.group(1)
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]
    return [vals[k] for k in order]


def sig_to_hex(sig, theme):
    """Couleur Excel (rgb ARGB ou (theme, teinte)) -> '#rrggbb', avec la formule de teinte d'Excel."""
    if isinstance(sig, str):
        return "#" + sig[-6:].lower()
    _, idx, tint = sig
    h = theme[idx]
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    hh, ll, ss = colorsys.rgb_to_hls(r, g, b)
    ll = ll * (1 + tint) if tint < 0 else ll * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(hh, min(1, max(0, ll)), ss)
    return "#%02x%02x%02x" % tuple(round(v * 255) for v in (r, g, b))


def iso(d):
    return d.date().isoformat() if isinstance(d, dt.datetime) else d.isoformat()


wb_f = openpyxl.load_workbook(SRC)
wb_v = openpyxl.load_workbook(SRC, data_only=True)

years, place_names, place_colors = {}, [], {}
theme = theme_colors(wb_f)
for ws in wb_f:
    if not ws.title.isdigit():
        continue
    year = int(ws.title)
    wv = wb_v[ws.title]

    # --- légende des lieux (cellules colorées de la colonne AB) ---
    legend = {}
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 28)
        if isinstance(c.value, str) and c.value not in NOT_A_PLACE:
            sig = fill_sig(c)
            if sig is not None:
                legend[sig] = c.value
                if c.value not in place_names and c.value not in PLACE_AS_TYPE:
                    place_names.append(c.value)
                    place_colors[c.value] = sig_to_hex(sig, theme)

    # --- grille des jours ---
    days = {}
    for m in range(12):
        for d in range(1, calendar.monthrange(year, m + 1)[1] + 1):
            r = 2 + d
            date_cell, code_cell = ws.cell(r, 2 + 2 * m), ws.cell(r, 3 + 2 * m)
            code = code_cell.value if isinstance(code_cell.value, str) and not code_cell.value.startswith("=") else None
            if code:
                code = CODE_ALIASES.get(code.strip(), code.strip())
            place = legend.get(fill_sig(date_cell)) or legend.get(fill_sig(code_cell))
            if place in PLACE_AS_TYPE:
                code, place = code or PLACE_AS_TYPE[place], None  # un code déjà saisi est conservé
            if code or place:
                entry = {}
                if code:
                    entry["c"] = code
                if place:
                    entry["p"] = place
                days[dt.date(year, m + 1, d).isoformat()] = entry

    # --- droits, jours fériés, vacances scolaires ---
    quotas, holidays, school = {}, [], []
    for r in range(3, ws.max_row + 1):
        label, val = wv.cell(r, 28).value, wv.cell(r, 29).value
        if isinstance(label, str) and r <= 14:
            code = CODE_ALIASES.get(label, label)
            if code in QUOTA_CODES and isinstance(val, (int, float)):
                quotas[code] = round(val, 6)
        if isinstance(label, (dt.datetime, dt.date)):
            holidays.append(iso(label))
        s, e = wv.cell(r, 30).value, wv.cell(r, 31).value
        if isinstance(label, str) and isinstance(s, (dt.datetime, dt.date)) and isinstance(e, (dt.datetime, dt.date)):
            school.append({"name": label, "start": iso(s), "end": iso(e)})
    years[str(year)] = {
        "quotas": quotas,
        "holidays": sorted(holidays),
        "schoolHolidays": school,
        "days": dict(sorted(days.items())),
    }

# Deux lieux ne doivent pas avoir la même couleur : le suivant reçoit une couleur de secours.
FALLBACK = ["#b8a1d9", "#6bb5a8", "#7fa6e0", "#d99a6c", "#a3c76b", "#c9a0a0"]
used, places = set(), []
for name in place_names:
    color = place_colors[name]
    if color in used:
        color = next(c for c in FALLBACK if c not in used)
    used.add(color)
    places.append({"name": name, "color": color})

seed = {"version": 2, "places": places, "years": dict(sorted(years.items()))}
with open(DST, "w", encoding="utf-8") as f:
    f.write("// Généré par tools/import_excel.py à partir de AF_Activites.xlsx\n")
    f.write("window.CONGES_SEED = ")
    json.dump(seed, f, ensure_ascii=False, indent=1)
    f.write(";\n")

for y, v in sorted(years.items()):
    cnt = {}
    for e in v["days"].values():
        if "c" in e:
            cnt[e["c"]] = cnt.get(e["c"], 0) + 1
    print(y, "fériés:", len(v["holidays"]), "vac:", len(v["schoolHolidays"]), "lieux:",
          sum(1 for e in v["days"].values() if "p" in e), "codes:", cnt, "droits:", v["quotas"])
print("lieux:", places)
